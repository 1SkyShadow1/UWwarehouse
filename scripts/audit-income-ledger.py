"""Read-only workbook/ledger reconciliation. Never writes to the source or ledger.

Usage: python audit-income-ledger.py --state STATE --output OUTPUT BOOK [BOOK ...]
Requires openpyxl for reading workbook values and original formulas.
"""
import argparse, collections, csv, datetime, hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook

def norm(value):
    return re.sub(r'[^a-z0-9]', '', str(value or '').lower())

def number(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None

def total(invoice):
    return sum(float(line.get('qty') or 0)*float(line.get('price') or 0) for line in invoice.get('items', []))

def audit(state_path, books, output):
    output.mkdir(parents=True, exist_ok=True)
    raw_state=state_path.read_bytes(); snapshot=json.loads(raw_state); db=snapshot['data']
    invoices=db.get('invoices', []); rows=[]; sources=[]; exceptions=[]
    def flag(area, record, reason, detail=''):
        exceptions.append(dict(area=area, record=record, reason=reason, detail=detail))
    seen_source={}
    for path in books:
        raw=load_workbook(path, data_only=False); cached=load_workbook(path, data_only=True)
        sources.append(dict(path=str(path), sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
        for sheet in raw:
            values=cached[sheet.title]
            for rr in sheet:
                for cell in rr:
                    if cell.data_type=='e' or values[cell.coordinate].data_type=='e':
                        flag('Workbook', f'{path.name} / {sheet.title}!{cell.coordinate}', 'Excel error', str(cell.value))
                    if cell.data_type=='f' and values[cell.coordinate].value is None:
                        flag('Workbook', f'{path.name} / {sheet.title}!{cell.coordinate}', 'Formula cache missing')
            for r in range(4, sheet.max_row+1):
                customer=values.cell(r, 2).value
                if not customer or str(customer).upper() in ['CUSTOMER','TOTAL','YTD']: continue
                get=lambda c:values[f'{c}{r}'].value
                date=get('A'); date=date.isoformat()[:10] if isinstance(date,(datetime.date,datetime.datetime)) else date
                row=dict(workbook=path.name, sheet=sheet.title, row=r, date=date, customer=customer, project=get('C'), number=get('D'))
                location=f'{path.name} / {sheet.title}!{r}'
                for col,key in [('E','quoted'),('F','deposit'),('G','amountDue'),('H','totalPaid'),('J','totalOwed'),('Q','currentOwed')]:
                    row[key]=number(get(col));row[key+'Formula']=sheet[f'{col}{r}'].value if sheet[f'{col}{r}'].data_type=='f' else None
                row['toPay']=get('I'); row['comments']=get('M');row['funding']=norm(row['number'])=='dt'
                fingerprint=tuple(str(row[k]) for k in ['date','customer','project','number','quoted','deposit','totalPaid'])
                row['duplicateOf']=seen_source.get(fingerprint, '')
                if row['duplicateOf']:flag('Income source',location,'Repeated source row',row['duplicateOf'])
                else:seen_source[fingerprint]=location
                matches=[]
                if row['number'] not in [None,'BB','SS','DT','']:
                    matches=[i for i in invoices if norm(i.get('id'))==norm(row['number']) and norm(i.get('customer'))==norm(customer)]
                if not matches:
                    matches=[i for i in invoices if norm(i.get('customer'))==norm(customer) and norm(i.get('project'))==norm(row['project']) and i.get('sourceSheet')==sheet.title and str(i.get('sourceWorkbook','')).endswith(path.name)]
                row['ledgerIds']='; '.join(dict.fromkeys(str(i['id']) for i in matches));row['matchCount']=len(matches)
                row['ledgerTotal']=total(matches[0]) if matches else None
                row['ledgerCollected']=float(matches[0].get('paid') or 0)+float(matches[0].get('deposit') or 0) if matches else None
                row['paymentConflict']=bool(re.search(r'to\s*pay',str(row['toPay'] or ''),re.I) or (row['totalOwed'] or 0)>0 or (row['currentOwed'] or 0)>0) and (row['totalPaid'] or 0)>0 and (row['totalPaid'] or 0)>=(row['quoted'] or 0)
                row['unpriced']=row['quoted'] is None or row['quoted']==0
                if not row['funding'] and not row['duplicateOf']:
                    if not matches:flag('Income source',location,'Not matched to invoice ledger',str(row['number'] or '')+' / '+str(customer))
                    if row['paymentConflict']:flag('Income source',location,'Paid amount conflicts with outstanding/To Pay','H='+str(row['totalPaid'])+' J='+str(row['totalOwed'])+' Q='+str(row['currentOwed']))
                    if row['unpriced']:flag('Income source',location,'No confirmed positive job price',str(row['quoted']))
                    if matches and row['quoted'] is not None and abs(row['quoted']-row['ledgerTotal'])>.02:flag('Income source',location,'Quote amount differs from ledger',str(row['quoted'])+' vs '+str(row['ledgerTotal']))
                    if matches and row['totalPaid'] is not None and abs(row['totalPaid']-row['ledgerCollected'])>.02:flag('Income source',location,'Paid amount differs from ledger',str(row['totalPaid'])+' vs '+str(row['ledgerCollected']))
                    if matches and date and str(matches[0].get('date'))!=str(date):flag('Income source',location,'Source date differs from ledger',str(date)+' vs '+str(matches[0].get('date')))
                    identifier=re.fullmatch(r'(?:BB|SS)(\d{2})(\d{2})(\d{2})\d{2}',str(row['number'] or ''))
                    if identifier and date:
                        try:
                            numbered_date=datetime.date(2000+int(identifier[3]),int(identifier[2]),int(identifier[1])).isoformat()
                            if numbered_date!=date:flag('Income source',location,'Source date differs from historical document number',str(date)+' vs '+numbered_date)
                        except ValueError:pass
                rows.append(row)
    grouped=collections.defaultdict(list)
    for invoice in invoices:grouped[str(invoice.get('id',''))].append(invoice)
    unique=[group[0] for key,group in grouped.items() if key]
    duplicates=[key for key,group in grouped.items() if len(group)>1]
    for key in duplicates:
        group=grouped[key]
        signatures={(norm(i.get('customer')),round(total(i),2),float(i.get('paid') or 0),float(i.get('deposit') or 0),i.get('date'),i.get('fy'),i.get('status')) for i in group}
        if len(signatures)>1:flag('Invoice',key,'Conflicting copies sharing an ID',f'{len(group)} copies / {len(signatures)} distinct financial/customer/date/status versions')
    for invoice in unique:
        key=invoice['id']; collected=float(invoice.get('paid') or 0)+float(invoice.get('deposit') or 0)
        if collected>total(invoice)+.02:flag('Invoice',key,'Collected exceeds invoice total',f'{collected} vs {total(invoice)}')
        if not invoice.get('items') or any(l.get('price') is None or float(l.get('price') or 0)<=0 for l in invoice.get('items',[])):flag('Invoice',key,'Missing/zero line price')
        if not re.fullmatch(r'\d{4}-\d{2}-\d{2}',str(invoice.get('date',''))):flag('Invoice',key,'Missing/invalid date')
    jobs=db.get('jobs',[]); linked={str(j.get('invoiceId','')) for j in jobs}
    for i in unique:
        if str(i['id']) not in linked:flag('Jobs',i['id'],'Invoice has no linked saved job')
    for j in jobs:
        if j.get('invoiceId') and str(j['invoiceId']) not in grouped:flag('Jobs',j.get('id'),'Broken invoice link',str(j['invoiceId']))
    duplicate_jobs=collections.Counter(str(j.get('id')) for j in jobs)
    for key,count in duplicate_jobs.items():
        if count>1:flag('Jobs',key,'Repeated job ID',str(count)+' jobs; ID-only edits are ambiguous')
    receipt_sums=collections.defaultdict(float)
    for receipt in db.get('receipts',[]):
        receipt_sums[str(receipt.get('invoiceId',''))]+=float(receipt.get('amount') or 0)
        if str(receipt.get('invoiceId','')) not in grouped:flag('Receipt',receipt.get('id'),'Broken invoice link',str(receipt.get('invoiceId')))
    for key,amount in receipt_sums.items():
        if key in grouped:
            i=grouped[key][0]; collected=float(i.get('paid') or 0)+float(i.get('deposit') or 0)
            if abs(amount-collected)>.02:flag('Receipt',key,'Receipt register differs from invoice collected amount',f'{amount} vs {collected}')
    for statement in db.get('fnbStatements',[]):
        transactions=statement.get('transactions',[]); amounts=[float(t.get('amount') or (-float(t.get('debit') or 0) if t.get('debit') else float(t.get('credit') or 0))) for t in transactions]
        credits=sum(a for a in amounts if a>0);debits=-sum(a for a in amounts if a<0)
        for name,actual in [('credits',credits),('debits',debits)]:
            if abs(actual-float(statement.get(name) or 0))>.02:flag('Bank',statement['id'],name+' differs from transaction detail',f'{actual} vs {statement.get(name)}')
        expected=float(statement.get('openingBalance') or 0)+credits-debits
        if abs(expected-float(statement.get('closingBalance') or 0))>.02:flag('Bank',statement['id'],'Opening + credits - debits differs from closing',f'{expected} vs {statement.get("closingBalance")}')
    for q in db.get('quotes',[]):
        if not q.get('items') or any(l.get('price') is None or float(l.get('price') or 0)<=0 for l in q.get('items',[])):flag('Quote',q.get('id'),'Missing/zero line price')
        text=json.dumps(q)
        if 'Imported quote' in text or re.search(r'"(?:introduction|notes)": "Imported from',text):flag('Quote',q.get('id'),'Import placeholder still present')
    for payable in db.get('payables',[]):
        if not float(payable.get('amount') or 0):flag('Payable',payable.get('id'),'No confirmed amount',str(payable.get('status')))
    for expense in db.get('expenses',[]):
        if not expense.get('date') or not float(expense.get('amount') or 0)>0:flag('Expense',expense.get('id') or expense.get('sourceId') or expense.get('desc'),'Missing date/non-positive amount')
    wage_total=sum((sum(float(w.get('owed') or 0) for w in wage.get('weeks',[])) or float(wage.get('monthly') or 0))-float(wage.get('advances') or 0) for wage in db.get('wages',[]))
    summary=dict(auditedAt=datetime.datetime.now(datetime.timezone.utc).isoformat(),revision=snapshot.get('revision'),ledgerSHA256=hashlib.sha256(raw_state).hexdigest(),sources=sources,collections={k:len(v) for k,v in db.items() if isinstance(v,list)},uniqueInvoices=len(unique),duplicateInvoiceCopies=len(invoices)-len(unique),invoiceTotal=round(sum(total(i) for i in unique),2),collected=round(sum(float(i.get('paid') or 0)+float(i.get('deposit') or 0) for i in unique),2),outstanding=round(sum(max(0,total(i)-float(i.get('paid') or 0)-float(i.get('deposit') or 0)) for i in unique),2),undatedPayrollObligation=wage_total,exceptionCounts=dict(collections.Counter(e['reason'] for e in exceptions)))
    (output/'financial-audit.json').write_text(json.dumps(dict(summary=summary,rows=rows,exceptions=exceptions),indent=2,default=str),encoding='utf-8')
    for name,records in [('income-reconciliation.csv',rows),('financial-exceptions.csv',exceptions)]:
        with (output/name).open('w',newline='',encoding='utf-8-sig') as f:
            writer=csv.DictWriter(f,fieldnames=list(records[0]) if records else ['area']);writer.writeheader();writer.writerows(records)
    print(json.dumps(summary,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--state',required=True,type=Path);parser.add_argument('--output',required=True,type=Path);parser.add_argument('books',nargs='+',type=Path);args=parser.parse_args()
    audit(args.state,args.books,args.output)
