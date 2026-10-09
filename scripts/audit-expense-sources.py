"""Read-only recursive evidence inventory and expense-row reconciliation.

No amounts are posted, no duplicates deleted and no source files modified.
Image-only documents are explicitly unresolved; text extraction is not OCR.
"""
import argparse, csv, hashlib, json, re, warnings
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document

def norm(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()

def digest(path):
    h = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()

def money(value):
    return round(float(value), 2) if isinstance(value, (int, float)) and not isinstance(value, bool) else None

def write_csv(path, rows, fields):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction='ignore')
        writer.writeheader(); writer.writerows(rows)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--state', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('roots', nargs='+')
    args = parser.parse_args()
    out = Path(args.output); out.mkdir(parents=True, exist_ok=True)
    snapshot = json.loads(Path(args.state).read_text(encoding='utf-8-sig')); db = snapshot['data']
    scans = {s.get('hash'): s for s in db.get('scannedDocuments', []) if s.get('hash')}
    inventory, rows, issues, representatives = [], [], [], {}
    for root in args.roots:
        for path in sorted(Path(root).rglob('*')):
            if not path.is_file(): continue
            entry = dict(path=str(path), size=path.stat().st_size, extension=path.suffix.lower())
            try:
                entry['sha256'] = digest(path)
                prior = representatives.get(entry['sha256'])
                entry['same_bytes_as'] = str(prior) if prior else ''
                if not prior: representatives[entry['sha256']] = path
            except Exception as error:
                entry['error'] = str(error)
            inventory.append(entry)
    print(f'Inventoried {len(inventory)} files; {len(representatives)} distinct hashes', flush=True)
    by_hash = {e.get('sha256'): e for e in inventory if not e.get('same_bytes_as')}
    for index, (hash_value, path) in enumerate(representatives.items()):
        entry = by_hash[hash_value]; ext = path.suffix.lower()
        scan = scans.get(hash_value)
        if scan:
            entry['scan_review'] = scan.get('reviewStatus', '')
            entry['scan_approved'] = scan.get('reviewStatus') == 'Approved' and scan.get('includedInTotals') is True
        try:
            if path.name.startswith('~$'):
                entry['inspection'] = 'Office owner lock file: no financial data'
                continue
            if ext == '.xlsx':
                with warnings.catch_warnings():
                    warnings.simplefilter('ignore')
                    book = load_workbook(path, read_only=True, data_only=True)
                entry['inspection'] = 'Workbook values inspected'
                for sheet in book:
                    header = None
                    for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                        normalized = [norm(v) for v in values]
                        if 'amount paid' in normalized and 'date' in normalized and 'description' in normalized:
                            header = {name: normalized.index(name) for name in ['date','description','amount paid']}
                            for name in ['customer','comment','invoice no','paid by']:
                                if name in normalized: header[name] = normalized.index(name)
                            continue
                        if not header: continue
                        get = lambda name: values[header[name]] if name in header and header[name] < len(values) else None
                        amount, when = money(get('amount paid')), get('date')
                        if not isinstance(when, (date, datetime)) or amount is None: continue
                        rows.append(dict(path=str(path), sha256=hash_value, sheet=sheet.title, row=number,
                            date=when.strftime('%Y-%m-%d'), vendor=str(get('customer') or ''),
                            description=str(get('description') or ''), source_category=str(get('comment') or ''),
                            invoice=str(get('invoice no') or ''), paid_by=str(get('paid by') or ''), amount=amount))
                book.close()
            elif ext == '.pdf':
                reader = PdfReader(path)
                texts = [page.extract_text() or '' for page in reader.pages]
                entry['pages'] = len(texts); entry['text_characters'] = sum(len(t.strip()) for t in texts)
                entry['inspection'] = 'Text extracted; financial interpretation pending' if entry['text_characters'] else 'Image-only PDF: visual/Gemini review required'
                entry['pages_without_text'] = sum(not t.strip() for t in texts)
            elif ext == '.docx':
                document = Document(path)
                text = '\n'.join([p.text for p in document.paragraphs] + [' | '.join(c.text for c in row.cells) for table in document.tables for row in table.rows])
                entry['text_characters'] = len(text.strip()); entry['inspection'] = 'Word paragraphs and tables extracted; financial interpretation pending'
            elif ext in ['.jpg','.jpeg','.png','.avif']:
                entry['inspection'] = 'Existing Gemini evidence located' if scan else 'Image: visual/Gemini review required'
            else:
                entry['inspection'] = 'Inventoried; format not parsed'
        except Exception as error:
            entry['inspection'] = 'Extraction failed'; entry['error'] = str(error)
        if index % 100 == 0: print(f'Inspected {index + 1}/{len(representatives)} distinct files', flush=True)
    seen_rows = {}; ledger = db.get('expenses', [])
    for row in rows:
        key = (row['date'],norm(row['vendor']),norm(row['description']),row['amount'],norm(row['invoice']))
        if key in seen_rows:
            row['status'] = 'Repeated source row: not added again'; row['duplicate_of'] = seen_rows[key]
            continue
        seen_rows[key] = f"{row['path']} | {row['sheet']} | {row['row']}"
        candidates = [e for e in ledger if e.get('date') == row['date'] and abs(float(e.get('amount') or 0)-row['amount']) < .02 and norm(e.get('vendor')) == norm(row['vendor'])]
        strong = [e for e in candidates if (row['invoice'] and norm(e.get('invoiceNo')) == norm(row['invoice'])) or norm(e.get('desc')) == norm(row['description'])]
        row['ledger_ids'] = '|'.join(str(e.get('id')) for e in strong)
        row['ledger_categories'] = '|'.join(sorted(set(e.get('category','') for e in strong)))
        row['status'] = 'Matched identity and amount' if len(strong) == 1 else 'Ambiguous ledger matches' if strong else 'Unmatched source row'
        if row['status'] != 'Matched identity and amount': issues.append(row.copy())
        if norm(row['description']) in ['material','materials'] and any(e.get('category') == 'Other' for e in strong):
            issue = row.copy(); issue['status'] = 'Material description categorized Other: category review required'; issues.append(issue)
    # Reproduce category panel using the application's exact source-row dedup key.
    distinct = {}; categories = defaultdict(float); dated = defaultdict(lambda: defaultdict(float))
    for e in ledger:
        key = (e.get('date'),norm(e.get('vendor')),norm(e.get('desc')),float(e.get('amount') or 0),norm(e.get('invoiceNo')),norm(e.get('fundedBy')))
        if key in distinct: continue
        distinct[key] = e; category = e.get('category') or 'Other'; amount = float(e.get('amount') or 0)
        categories[category] += amount; dated[str(e.get('date',''))[:4]][category] += amount
    fields = ['path','sha256','sheet','row','date','vendor','description','source_category','invoice','paid_by','amount','status','ledger_ids','ledger_categories','duplicate_of']
    write_csv(out/'expense-source-rows.csv', rows, fields)
    write_csv(out/'expense-source-exceptions.csv', issues, fields)
    write_csv(out/'source-file-inventory.csv', inventory, ['path','size','extension','sha256','same_bytes_as','inspection','pages','pages_without_text','text_characters','scan_review','scan_approved','error'])
    summary = dict(revision=snapshot.get('revision'),file_count=len(inventory),unique_file_hashes=len(representatives),
        duplicate_file_copies=len(inventory)-len(representatives),source_expense_rows=len(rows),
        row_status_counts=dict(Counter(r['status'] for r in rows)),exception_counts=dict(Counter(r['status'] for r in issues)),
        inspection_counts=dict(Counter(e.get('inspection','Duplicate byte copy') for e in inventory)),
        categories={k:round(v,2) for k,v in categories.items()},categories_by_transaction_year={y:{k:round(v,2) for k,v in cat.items()} for y,cat in dated.items()},
        undated_payroll_estimate=round(sum((sum(float(week.get('owed') or 0) for week in w.get('weeks',[])) or float(w.get('monthly') or 0))-float(w.get('advances') or 0) for w in db.get('wages',[])),2),notes=['Full recursive inventory is not proof of payment or semantic verification of all documents.','Image-only, unsupported and failed sources require review; amounts are not guessed.','No source or ledger records changed. Exact duplicate bytes are inventoried once for extraction, every path retained.','Payroll estimate is separate from recorded expense categories.'])
    (out/'expense-source-audit.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps(summary,indent=2), flush=True)

if __name__ == '__main__': main()
